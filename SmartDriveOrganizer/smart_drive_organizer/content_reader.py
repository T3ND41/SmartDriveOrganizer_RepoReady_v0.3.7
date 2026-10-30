from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Optional

MAX_READ_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 12000


def _plain_text(path: Path) -> str:
    data = path.read_bytes()[:MAX_READ_BYTES]
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding, errors="ignore")[:MAX_TEXT_CHARS]
        except (LookupError, UnicodeError):
            continue
    return ""


def _docx_text(path: Path) -> str:
    from xml.etree import ElementTree as ET
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    root = ET.fromstring(xml)
    texts = [node.text or "" for node in root.iter() if node.tag.endswith("}t")]
    return " ".join(texts)[:MAX_TEXT_CHARS]


def _pdf_text(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    pieces = []
    for page in reader.pages[:8]:
        pieces.append(page.extract_text() or "")
        if sum(map(len, pieces)) >= MAX_TEXT_CHARS:
            break
    return " ".join(pieces)[:MAX_TEXT_CHARS]


def _xlsx_text(path: Path) -> str:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    pieces = []
    try:
        for sheet in workbook.worksheets[:3]:
            for row in sheet.iter_rows(max_row=80, max_col=20, values_only=True):
                values = [str(value) for value in row if value is not None]
                if values:
                    pieces.append(" ".join(values))
                if sum(map(len, pieces)) >= MAX_TEXT_CHARS:
                    break
            if sum(map(len, pieces)) >= MAX_TEXT_CHARS:
                break
    finally:
        workbook.close()
    return " ".join(pieces)[:MAX_TEXT_CHARS]


def _pptx_text(path: Path) -> str:
    from pptx import Presentation
    presentation = Presentation(str(path))
    pieces = []
    for slide in presentation.slides[:15]:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                pieces.append(shape.text)
            if sum(map(len, pieces)) >= MAX_TEXT_CHARS:
                break
    return " ".join(pieces)[:MAX_TEXT_CHARS]


def read_document_text(path: Path) -> str:
    """Best-effort local text extraction. Missing optional packages simply skip that format."""
    suffix = path.suffix.lower()
    try:
        if suffix in {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".htm", ".log", ".yaml", ".yml", ".sql", ".tex"}:
            text = _plain_text(path)
            if suffix in {".html", ".htm", ".xml"}:
                text = re.sub(r"<[^>]+>", " ", text)
            return text
        if suffix == ".rtf":
            text = _plain_text(path)
            return re.sub(r"\\[a-zA-Z]+-?\d* ?|[{}]", " ", text)[:MAX_TEXT_CHARS]
        if suffix == ".docx":
            return _docx_text(path)
        if suffix == ".pdf":
            return _pdf_text(path)
        if suffix in {".xlsx", ".xlsm"}:
            return _xlsx_text(path)
        if suffix == ".pptx":
            return _pptx_text(path)
    except (ImportError, OSError, ValueError, KeyError, zipfile.BadZipFile, RuntimeError):
        return ""
    return ""
